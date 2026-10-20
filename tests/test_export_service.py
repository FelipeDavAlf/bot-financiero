from datetime import date
import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from core.database.session import Base
from core.database.models import TransactionRecord
from repositories.user_repository import UserRepository
from repositories.transaction_repository import TransactionRepository
from services.export_service import ExcelExportService


@pytest.fixture
def db_session():
    """Provides an in-memory database session for testing."""
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()


def test_excel_export_service_generates_valid_workbook():
    service = ExcelExportService()

    # Mock transactions
    tx1 = TransactionRecord(
        id=1,
        user_id=1,
        date=date(2026, 10, 1),
        transaction_type="Ingreso",
        amount=15000.0,
        source_account_name="Empresa",
        destination_account_name="Tarjeta Débito BBVA",
        category="Salario",
        subcategory="Nómina",
        description="Pago de nómina quincenal",
    )
    tx2 = TransactionRecord(
        id=2,
        user_id=1,
        date=date(2026, 10, 2),
        transaction_type="Gasto",
        amount=350.0,
        source_account_name="Tarjeta Débito BBVA",
        destination_account_name=None,
        category="Alimentación",
        subcategory="Supermercado",
        description="Despensa Walmart",
    )
    tx3 = TransactionRecord(
        id=3,
        user_id=1,
        date=date(2026, 10, 3),
        transaction_type="Traspaso",
        amount=1000.0,
        source_account_name="Tarjeta Débito BBVA",
        destination_account_name="Cuenta Mercado Pago",
        category="Traspaso",
        subcategory="Ahorro",
        description="Ahorro a Mercado Pago",
    )

    transactions = [tx1, tx2, tx3]
    buffer = service.generate_user_finance_workbook(user_name="Felipe", transactions=transactions)

    assert buffer is not None
    assert buffer.getbuffer().nbytes > 0

    # Load workbook from generated buffer and verify content
    wb = openpyxl.load_workbook(buffer)
    assert "Resumen Financiero" in wb.sheetnames
    assert "Presupuesto y Metas" in wb.sheetnames
    assert "Cuentas y Métodos" in wb.sheetnames
    assert "Detalle de Movimientos" in wb.sheetnames

    # Check Summary Sheet
    ws_summary = wb["Resumen Financiero"]
    assert "REPORTE DE FINANZAS PERSONALES" in str(ws_summary["B2"].value)
    assert ws_summary["B6"].value == 15000.0  # Total Ingresos
    assert ws_summary["C6"].value == 350.0    # Total Gastos
    assert ws_summary["D6"].value == 14650.0  # Balance Neto
    assert ws_summary["E6"].value == 1000.0   # Traspasos
    assert ws_summary["F6"].value is not None  # Tasa de ahorro
    assert ws_summary["G6"].value is not None  # Estado financiero

    # Check Budget Sheet
    ws_budget = wb["Presupuesto y Metas"]
    assert "PLANIFICADOR" in str(ws_budget["B2"].value)
    assert ws_budget["B5"].value == "Categoría"
    assert ws_budget["D5"].value == "Presupuesto Asignado (Meta)"
    assert ws_budget["E5"].value == "Gasto Real Actual"

    # Check Accounts Sheet
    ws_accounts = wb["Cuentas y Métodos"]
    assert "ANÁLISIS DE CUENTAS" in str(ws_accounts["B2"].value)
    assert ws_accounts["B5"].value == "Cuenta / Tarjeta"
    assert ws_accounts["C5"].value == "Tipo de Cuenta"

    # Check Detail Sheet
    ws_detail = wb["Detalle de Movimientos"]
    assert ws_detail["A1"].value == "#"
    assert ws_detail["B1"].value == "Fecha"
    assert ws_detail["D1"].value == "Monto"
    assert ws_detail.max_row == 4  # Header + 3 transactions


def test_excel_export_service_handles_empty_transactions():
    service = ExcelExportService()
    buffer = service.generate_user_finance_workbook(user_name="Usuario Vacio", transactions=[])

    assert buffer is not None
    wb = openpyxl.load_workbook(buffer)
    assert "Resumen Financiero" in wb.sheetnames
    assert "Presupuesto y Metas" in wb.sheetnames

    ws_summary = wb["Resumen Financiero"]
    assert ws_summary["B6"].value == 0.0
    assert ws_summary["C6"].value == 0.0
    assert ws_summary["D6"].value == 0.0



def test_transaction_repository_get_all_user_transactions(db_session):
    user_repo = UserRepository(db_session)
    user = user_repo.create_user(telegram_id=98765, first_name="Test User")

    tx_repo = TransactionRepository(db_session)

    # Insert 2 transactions
    record1 = TransactionRecord(
        user_id=user.id,
        date=date(2026, 10, 5),
        transaction_type="Gasto",
        amount=50.0,
        category="Alimentación",
        subcategory="Café",
        description="Café americano",
    )
    record2 = TransactionRecord(
        user_id=user.id,
        date=date(2026, 10, 6),
        transaction_type="Ingreso",
        amount=2000.0,
        category="Salario",
        subcategory="Bono",
        description="Bono puntualidad",
    )
    db_session.add(record1)
    db_session.add(record2)
    db_session.commit()

    all_txs = tx_repo.get_all_user_transactions(user_id=user.id)
    assert len(all_txs) == 2
    assert all_txs[0].amount == 50.0
    assert all_txs[1].amount == 2000.0
